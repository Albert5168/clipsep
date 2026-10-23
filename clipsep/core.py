"""Core pipeline: OCR text off an image, erase it, and describe editable text blocks."""
from __future__ import annotations

import io
from dataclasses import dataclass, field, asdict

import cv2
import numpy as np

_engine = None


def get_engine():
    global _engine
    if _engine is None:
        from rapidocr import RapidOCR
        _engine = RapidOCR(params={"Global.log_level": "error"})
    return _engine


@dataclass
class Run:
    text: str
    color: tuple[int, int, int]  # RGB


@dataclass
class Line:
    x0: int
    y0: int
    x1: int
    y1: int
    runs: list[Run]
    em: float = 0.0

    @property
    def h(self):
        return self.y1 - self.y0

    @property
    def text(self):
        return "".join(r.text for r in self.runs)


@dataclass
class Block:
    x: int
    y: int
    w: int
    h: int
    lines: list[list[Run]]
    font_px: float
    bold: bool = False
    font: str = "Microsoft JhengHei"
    align: str = "left"
    indents: list[float] = field(default_factory=list)  # px from block left, per line

    @property
    def text(self):
        return "\n".join("".join(r.text for r in ln) for ln in self.lines)

    def to_dict(self):
        return asdict(self)

    @staticmethod
    def from_dict(d):
        d = dict(d)
        d["lines"] = [[Run(r["text"], tuple(r["color"])) for r in ln] for ln in d["lines"]]
        return Block(**d)


@dataclass
class Page:
    width: int
    height: int
    original: np.ndarray  # BGR
    background: np.ndarray  # BGR, text removed
    blocks: list[Block] = field(default_factory=list)


# ---------- colour helpers ----------

def _two_colors(pixels: np.ndarray):
    """Split pixels into (text colour, background colour) using 2-means; text = minority."""
    data = pixels.reshape(-1, 3).astype(np.float32)
    if len(data) < 4:
        c = data.mean(0)
        return c, c
    crit = (cv2.TERM_CRITERIA_EPS + cv2.TERM_CRITERIA_MAX_ITER, 20, 1.0)
    _, labels, centers = cv2.kmeans(data, 2, None, crit, 3, cv2.KMEANS_PP_CENTERS)
    counts = np.bincount(labels.ravel(), minlength=2)
    t = int(np.argmin(counts))
    return centers[t], centers[1 - t]


def _bgr_to_rgb(c):
    return int(c[2]), int(c[1]), int(c[0])


def _core_color(patch: np.ndarray, bg: np.ndarray) -> np.ndarray:
    """Colour of the glyph cores: the pixels farthest from the background (skips anti-aliasing)."""
    px = patch.reshape(-1, 3).astype(np.float32)
    d = np.linalg.norm(px - bg, axis=1)
    return np.median(px[d >= np.percentile(d, 88)], axis=0)


def _lab(rgb):
    return cv2.cvtColor(np.uint8([[rgb]]), cv2.COLOR_RGB2LAB)[0, 0].astype(np.float32)


def _merge_colors(colors: list[tuple], tol=22.0) -> list[tuple]:
    """Collapse near-identical per-character colours (Lab distance) to one shared colour each."""
    clusters: list[list[int]] = []
    labs = [_lab(c) for c in colors]
    for i, lab in enumerate(labs):
        for cl in clusters:
            if np.linalg.norm(labs[cl[0]] - lab) < tol:
                cl.append(i)
                break
        else:
            clusters.append([i])
    out = list(colors)
    for cl in clusters:
        mean = tuple(int(v) for v in np.median([colors[i] for i in cl], axis=0))
        for i in cl:
            out[i] = mean
    return out


# ---------- OCR + grouping ----------

def _ocr_lines(img: np.ndarray, min_score=0.3, drop=("Gemini Notebook", "NotebookLM")) -> list[Line]:
    res = get_engine()(img, return_word_box=True)
    if res.boxes is None:
        return []
    lines = []
    for box, txt, score, words in zip(res.boxes, res.txts, res.scores, res.word_results):
        if score < min_score or not txt.strip():
            continue
        if any(d.replace(" ", "") in txt.replace(" ", "") for d in drop):
            continue
        # text drawn at an angle inside artwork (and stray 1-2 char guesses) stays in the picture
        (tx0, ty0), (tx1, ty1) = box[0], box[1]
        if abs(np.degrees(np.arctan2(ty1 - ty0, max(tx1 - tx0, 1)))) > 4:
            continue
        if len(txt.strip()) <= 2 and score < 0.8:
            continue
        xs, ys = box[:, 0], box[:, 1]
        x0, y0, x1, y1 = int(xs.min()), int(ys.min()), int(xs.max()), int(ys.max())
        # one colour per character, then merge neighbours of the same colour into runs
        line_fg, line_bg = _two_colors(img[y0:y1, x0:x1])
        line_rgb = _bgr_to_rgb(_core_color(img[y0:y1, x0:x1], line_bg))
        chars, colors = [], []
        for ch, _s, wb in words:
            wb = np.array(wb)
            cx0, cy0 = max(int(wb[:, 0].min()), 0), max(int(wb[:, 1].min()), 0)
            cx1, cy1 = int(wb[:, 0].max()), int(wb[:, 1].max())
            patch = img[cy0:cy1, cx0:cx1]
            chars.append(ch)
            colors.append(_bgr_to_rgb(_core_color(patch, line_bg)) if patch.size else line_rgb)
        # punctuation has too few pixels to measure; it takes the colour of the text before it
        for i, ch in enumerate(chars):
            if not ch.isalnum():
                colors[i] = colors[i - 1] if i else next(
                    (c for c, h in zip(colors, chars) if h.isalnum()), line_rgb)
        runs: list[Run] = []
        for ch, rgb in zip(chars, _merge_colors(colors)):
            if runs and runs[-1].color == rgb:
                runs[-1].text += ch
            else:
                runs.append(Run(ch, rgb))
        # OCR word boxes lose spaces; fall back to whole text when the runs don't reproduce it
        if "".join(r.text for r in runs) != txt:
            if len(runs) == 1 or not runs:
                runs = [Run(txt, line_rgb)]
            else:
                runs = _respace(runs, txt)
        # font size from line width: detection boxes are padded vertically, not horizontally
        em = (x1 - x0) / max(sum(_char_em(c) for c in txt), 1)
        lines.append(Line(x0, y0, x1, y1, runs, em))
    return lines


def _char_em(c: str) -> float:
    """Approximate advance width of a character in ems."""
    if c == " ":
        return 0.3
    if ord(c) >= 0x2E80 or c in "（）：，。、；！？":
        return 1.0
    if c in "iIl.,:;'|!":
        return 0.3
    return 0.58


def _respace(runs: list[Run], txt: str) -> list[Run]:
    """Re-insert the spaces from the full OCR string into colour runs."""
    out, i = [], 0
    for r in runs:
        buf = ""
        for ch in r.text:
            while i < len(txt) and txt[i] != ch:
                buf += txt[i]
                i += 1
            buf += ch
            i += 1
        out.append(Run(buf, r.color))
    if i < len(txt):
        out[-1].text += txt[i:]
    return out


def _group(lines: list[Line]) -> list[list[Line]]:
    """Merge consecutive lines of a paragraph: same left edge, similar height, small gap."""
    lines = sorted(lines, key=lambda l: (l.y0, l.x0))
    groups: list[list[Line]] = []
    for ln in lines:
        placed = False
        for g in groups:
            last = g[-1]
            # same left edge, or a hanging indent ("1. ..." then indented continuation lines)
            same_left = -max(10, 0.5 * last.h) <= ln.x0 - g[0].x0 <= max(10, 3 * last.em)
            similar_h = abs(ln.h - last.h) <= 0.3 * last.h
            gap = ln.y0 - last.y1
            if same_left and similar_h and -0.2 * last.h <= gap <= 0.8 * last.h:
                g.append(ln)
                placed = True
                break
        if not placed:
            groups.append([ln])
    return groups


def _is_bold(img, line: Line) -> bool:
    """Stroke thickness vs line height: thick strokes read as bold."""
    gray = cv2.cvtColor(img[line.y0:line.y1, line.x0:line.x1], cv2.COLOR_BGR2GRAY)
    if gray.size == 0:
        return False
    _, bw = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
    if (bw > 0).mean() > 0.5:
        bw = 255 - bw
    dist = cv2.distanceTransform(bw, cv2.DIST_L2, 3)
    strokes = dist[dist > 0]
    if strokes.size == 0:
        return False
    return float(np.percentile(strokes, 90)) * 2 / line.h > 0.11


# ---------- text removal ----------

def _text_mask(img: np.ndarray, lines: list[Line]):
    """Mask of text pixels (incl. anti-aliased edges) and a flat-fill layer for uniform backgrounds."""
    mask = np.zeros(img.shape[:2], np.uint8)
    flat = img.copy()
    flat_mask = np.zeros(img.shape[:2], np.uint8)
    for ln in lines:
        pad = max(3, ln.h // 6)
        y0, y1 = max(ln.y0 - pad, 0), min(ln.y1 + pad, img.shape[0])
        x0, x1 = max(ln.x0 - pad, 0), min(ln.x1 + pad, img.shape[1])
        patch = img[y0:y1, x0:x1].astype(np.float32)
        fg, bg = _two_colors(patch)
        contrast = max(float(np.linalg.norm(fg - bg)), 1.0)
        d_bg = np.linalg.norm(patch - bg, axis=2)
        m = (d_bg > 0.15 * contrast).astype(np.uint8) * 255
        k = max(5, ln.h // 5) | 1
        m = cv2.dilate(m, np.ones((k, k), np.uint8))
        mask[y0:y1, x0:x1] = np.maximum(mask[y0:y1, x0:x1], m)
        # background pixels around the text: if they barely vary, fill flat instead of inpainting
        bg_px = patch[m == 0]
        if len(bg_px) > 20 and float(bg_px.std(0).max()) < 10:
            fill = np.median(bg_px, 0)
            sub = flat[y0:y1, x0:x1]
            sub[m > 0] = fill
            flat_mask[y0:y1, x0:x1] = np.maximum(flat_mask[y0:y1, x0:x1], m)
    return mask, flat, flat_mask


def erase_text(img: np.ndarray, lines: list[Line], drop_boxes=()) -> np.ndarray:
    mask, flat, flat_mask = _text_mask(img, lines)
    for (x0, y0, x1, y1) in drop_boxes:
        mask[y0:y1, x0:x1] = 255
    rest = cv2.subtract(mask, flat_mask)
    out = cv2.inpaint(img, mask, 7, cv2.INPAINT_TELEA)
    out[flat_mask > 0] = flat[flat_mask > 0]
    if rest.any():
        out = cv2.inpaint(out, rest, 7, cv2.INPAINT_TELEA)
    return out


# ---------- public API ----------

def _unify_colors(lines: list[list[Run]]):
    """Same colour across a paragraph's lines, then join neighbouring runs that now match."""
    flat = [r for ln in lines for r in ln]
    for r, c in zip(flat, _merge_colors([r.color for r in flat])):
        r.color = c
    for ln in lines:
        merged: list[Run] = []
        for r in ln:
            if merged and merged[-1].color == r.color:
                merged[-1].text += r.text
            else:
                merged.append(r)
        ln[:] = merged


_s2t = None


def _to_traditional(text: str) -> str:
    """The OCR model sometimes returns simplified variants of traditional characters."""
    global _s2t
    if _s2t is None:
        from opencc import OpenCC
        _s2t = OpenCC("s2t")
    return _s2t.convert(text)


def analyze(img: np.ndarray, drop_watermark=True, traditional=True) -> Page:
    drop = ("Gemini Notebook", "NotebookLM") if drop_watermark else ()
    lines_all = _ocr_lines(img, drop=())
    if traditional:
        for ln in lines_all:
            for r in ln.runs:
                r.text = _to_traditional(r.text)
    keep, gone = [], []
    for ln in lines_all:
        (gone if any(d.replace(" ", "") in ln.text.replace(" ", "") for d in drop) else keep).append(ln)
    # watermark: also erase its small logo just left of the text
    drop_boxes = [(max(l.x0 - l.h * 2, 0), l.y0 - 2, l.x1 + 2, l.y1 + 2) for l in gone]
    background = erase_text(img, keep, drop_boxes)

    blocks = []
    for g in _group(keep):
        x0 = min(l.x0 for l in g)
        y0 = min(l.y0 for l in g)
        x1 = max(l.x1 for l in g)
        y1 = max(l.y1 for l in g)
        font_px = float(np.median([l.em for l in g]))
        bold = False
        _unify_colors([l.runs for l in g])
        blocks.append(Block(x0, y0, x1 - x0, y1 - y0, [l.runs for l in g], font_px, bold,
                            indents=[float(l.x0 - x0) for l in g]))
    # stroke width doesn't separate weights reliably; treat clearly larger text as headings
    if blocks:
        body = float(np.median([b.font_px for b in blocks]))
        for b in blocks:
            b.bold = b.font_px > 1.5 * body
    h, w = img.shape[:2]
    return Page(w, h, img, background, blocks)


def load_images(path: str) -> list[np.ndarray]:
    """Images from a .pptx (one per slide), PDF (one per page), or a single image file."""
    low = path.lower()
    if low.endswith(".pptx"):
        return _pptx_images(path)
    if low.endswith(".pdf"):
        import pymupdf
        doc = pymupdf.open(path)
        out = []
        for p in doc:
            pix = p.get_pixmap(dpi=200)
            arr = np.frombuffer(pix.samples, np.uint8).reshape(pix.h, pix.w, pix.n)
            out.append(cv2.cvtColor(arr, cv2.COLOR_RGB2BGR if pix.n == 3 else cv2.COLOR_RGBA2BGR))
        return out
    data = np.fromfile(path, np.uint8)  # handles non-ASCII Windows paths
    return [cv2.imdecode(data, cv2.IMREAD_COLOR)]


def _pptx_images(path: str) -> list[np.ndarray]:
    """Largest picture on each slide (these decks are one flattened image per slide)."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = Presentation(path)
    out = []
    for slide in prs.slides:
        pics = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
        if not pics:
            continue
        pic = max(pics, key=lambda s: s.width * s.height)
        arr = np.frombuffer(pic.image.blob, np.uint8)
        out.append(cv2.imdecode(arr, cv2.IMREAD_COLOR))
    return out


def encode_png(img: np.ndarray) -> bytes:
    ok, buf = cv2.imencode(".png", img)
    return buf.tobytes()
