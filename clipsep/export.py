"""Write pages to a .pptx: clean background picture + editable text boxes."""
from __future__ import annotations

import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from .core import Block, Page, encode_png

SLIDE_W = Emu(12192000)  # 16:9, 13.333 in
EMU_PER_PT = 12700
LINE_SPACING = 1.2


# Layout helpers, in image pixels. The editor uses the same ones so what you see is what exports.

def line_pitch(b: Block) -> float:
    n = len(b.lines)
    pitch = b.h / n if n > 1 else b.font_px * LINE_SPACING
    return max(pitch, b.font_px * 1.05)


def text_top(b: Block) -> float:
    """Detection boxes are taller than the glyphs; centre the first line in its box."""
    return b.y + (b.h / max(len(b.lines), 1) - line_pitch(b)) / 2


def box_width(b: Block) -> float:
    """A little slack so a wider font doesn't wrap the line."""
    return b.w + b.w * 0.08 + b.font_px * 1.5


def _set_font(run, name):
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", name)


def export_pptx(pages: list[Page], out_path: str):
    prs = Presentation()
    if pages:
        prs.slide_width = SLIDE_W
        prs.slide_height = Emu(int(SLIDE_W * pages[0].height / pages[0].width))
    blank = prs.slide_layouts[6]
    for page in pages:
        slide = prs.slides.add_slide(blank)
        # pages of another aspect ratio are fitted inside the slide, centred
        scale = min(prs.slide_width / page.width, prs.slide_height / page.height)
        ox = (prs.slide_width - page.width * scale) / 2
        oy = (prs.slide_height - page.height * scale) / 2
        slide.shapes.add_picture(io.BytesIO(encode_png(page.background)), int(ox), int(oy),
                                 int(page.width * scale), int(page.height * scale))
        for b in page.blocks:
            font_pt = b.font_px * scale / EMU_PER_PT
            tb = slide.shapes.add_textbox(int(ox + b.x * scale), int(oy + text_top(b) * scale),
                                          int(box_width(b) * scale),
                                          int(line_pitch(b) * len(b.lines) * scale))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.vertical_anchor = MSO_ANCHOR.TOP
            pitch_pt = line_pitch(b) * scale / EMU_PER_PT
            for i, line in enumerate(b.lines):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.alignment = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(b.align, PP_ALIGN.LEFT)
                p.line_spacing = Pt(pitch_pt)
                indent = b.indents[i] if i < len(b.indents) else 0
                if indent > 2:
                    p._p.get_or_add_pPr().set("marL", str(int(indent * scale)))
                for r in line:
                    run = p.add_run()
                    run.text = r.text
                    run.font.size = Pt(round(font_pt, 1))
                    run.font.bold = b.bold
                    run.font.color.rgb = RGBColor(*r.color)
                    _set_font(run, b.font)
    prs.save(out_path)
